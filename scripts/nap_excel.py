"""
Nạp dữ liệu từ file báo cáo tồn kho Excel (format BaoCao_TonKho_<Khoi>_dd_mm_yyyy.xlsx)
vào schema app_inventory trên Supabase.

Nạp gì:
  1. dm_hang_ton_kho  — master mặt hàng (đơn giá, BU, nhóm SP, PL hàng hóa...) từ sheet
                        DuLieu_goc. CÓ PHIÊN BẢN: chỉ thêm dòng mới cho mã mới hoặc mã có
                        thuộc tính thay đổi so với phiên bản gần nhất; mã không đổi giữ nguyên.
  2. dm_nhom_bao_cao  — thứ tự dòng nhóm SP theo sheet báo cáo.
  3. gv_ban_tb        — cột "GV bán trung bình tháng gần nhất" (kỳ = ngày báo cáo).
  4. ton_kho_ky       — số liệu kỳ báo cáo + các kỳ trước có trong file (xu hướng).

Cách chạy (cần Supabase CLI đã `supabase link`):
  python scripts/nap_excel.py --file "BaoCao_TonKho_NgoaiKhoa_30_09_2026.xlsx"
  python scripts/nap_excel.py --file ... --ky 2026-09-30 --sql-only out/   # chỉ sinh SQL

--ky: ngày báo cáo; mặc định đọc từ tên file (dd_mm_yyyy). Đây cũng là updated_at
      (ngày hiệu lực) của phiên bản master mới.
"""
import argparse
import hashlib
import os
import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

SHEET_RAW = "DuLieu_goc"
CHUNK = 800

# Cột trong sheet báo cáo (hàng 6 trở đi) -> chỉ tiêu
COT_BAO_CAO = {
    "C": "hsd_tren_12t", "D": "hsd_6_12t", "E": "hsd_duoi_6t",
    "G": "het_han", "H": "cho_huy", "I": "loi_kt",
    "K": "clc_tren_24t", "L": "clc_tren_12t", "M": "clc_tren_06t",
}
NHAN_KHONG_PHAI_NHOM = ("tổng", "tỷ trọng", "tăng/giảm", "% tăng/giảm")
# Bảng "Cảnh báo tồn đọng" mức khối
CANH_BAO = {
    "hàng hết date, hỏng": "het_han",
    "hàng đã duyệt chờ hủy": "cho_huy",
    "hàng lỗi/kt": "loi_kt",
    "tổng giá trị clc (chậm luân chuyển)": "clc",
    "tổng tồn kho (hàng đang kd)": "tong_kd",
}


def q(v):
    """Literal SQL."""
    if v is None:
        return "null"
    if isinstance(v, (int, float)):
        return repr(float(v)) if isinstance(v, float) else str(v)
    if isinstance(v, (date, datetime)):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def parse_ngay(s):
    return datetime.strptime(s.strip(), "%d/%m/%Y").date()


def ky_tu_ten_file(path):
    m = re.search(r"(\d{2})_(\d{2})_(\d{4})", Path(path).stem)
    if not m:
        sys.exit("Không đọc được ngày từ tên file, truyền --ky YYYY-MM-DD")
    return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def doc_master(wb, ten_nhom_chuan):
    """1 dòng / mã từ DuLieu_goc. Nhóm SP chuẩn hoá theo cách viết trên sheet báo cáo
    (SUMIFS của Excel không phân biệt hoa thường)."""
    ws = wb[SHEET_RAW]
    rows = ws.iter_rows(values_only=True)
    hdr = [norm(h) for h in next(rows)]
    ix = {h: i for i, h in enumerate(hdr)}
    need = ["Mã mặt hàng", "Tên mặt hàng", "Khối", "BU", "Nhóm SP", "Hãng", "PL hàng hóa", "Đơn giá 2"]
    for n in need:
        if n not in ix:
            sys.exit(f"Thiếu cột '{n}' trong {SHEET_RAW}")
    out = {}
    for r in rows:
        ma = norm(r[ix["Mã mặt hàng"]])
        if not ma or ma in out:
            continue
        nhom = norm(r[ix["Nhóm SP"]])
        nhom = ten_nhom_chuan.get(nhom.lower(), nhom)
        out[ma] = {
            "ma_bravo": ma,
            "ten_mat_hang": norm(r[ix["Tên mặt hàng"]]),
            "khoi": norm(r[ix["Khối"]]),
            "bu": norm(r[ix["BU"]]),
            "nhom_sp": nhom,
            "hang": norm(r[ix["Hãng"]]),
            "pl_hang_hoa": norm(r[ix["PL hàng hóa"]]),
            "don_gia": round(float(r[ix["Đơn giá 2"]] or 0), 6),
        }
    for v in out.values():
        key = "|".join(str(v[k]) for k in
                       ("ten_mat_hang", "khoi", "bu", "nhom_sp", "hang", "pl_hang_hoa", "don_gia"))
        v["row_hash"] = hashlib.md5(key.encode("utf-8")).hexdigest()
    return list(out.values())


def doc_bao_cao(wb_f, wb_v, khoi_du_lieu, ky):
    """Đọc sheet báo cáo (sheet thứ 2): thứ tự nhóm, GV bán TB, số liệu kỳ này + kỳ trước."""
    ten_sheet = [s for s in wb_v.sheetnames if s != SHEET_RAW][0]
    wf, wv = wb_f[ten_sheet], wb_v[ten_sheet]
    val = lambda c: wv[c].value
    # cột kỳ trước: V5.. (dd/mm/yyyy)
    ky_truoc = {}
    for col in "VWXYZ":
        v = val(f"{col}5")
        if isinstance(v, str) and re.match(r"\d{2}/\d{2}/\d{4}", v.strip()):
            ky_truoc[col] = parse_ngay(v)
    m = re.search(r"\((\d{2}/\d{2}/\d{4})", str(val("R4") or ""))
    ky_R = parse_ngay(m.group(1)) if m else None
    m = re.search(r"\((\d{2}/\d{2}/\d{4})\)", str(val("T4") or ""))
    ky_T = parse_ngay(m.group(1)) if m else None

    thu_tu, gv, tkk = [], [], []
    bu = None
    for r in range(6, wv.max_row + 1):
        a, b = norm(val(f"A{r}")), norm(val(f"B{r}"))
        if a.upper().startswith("BU "):
            bu = a[3:].strip()
        elif a.upper().startswith("KHỐI"):
            bu = None
            continue
        if not bu or not b:
            if a.upper().startswith("CẢNH BÁO"):
                break
            continue
        bl = b.lower()
        if bl == "tổng":
            # mức BU: kỳ trước (R), dang_sd 2 kỳ trước (T), xấu 2 kỳ trước (U)
            if ky_R and isinstance(val(f"R{r}"), (int, float)):
                tkk.append((ky_R, bu, "", "tong_kd", val(f"R{r}")))
            for col, ct in (("T", "dang_sd"), ("U", "xau")):
                f = str(wf[f"{col}{r}"].value or "")
                mm = re.search(r"-\s*([\d.]+)\)\s*/", f)
                if ky_T and mm:
                    tkk.append((ky_T, bu, "", ct, float(mm.group(1))))
            continue
        if any(bl.startswith(x) for x in NHAN_KHONG_PHAI_NHOM):
            continue
        # dòng nhóm SP
        thu_tu.append((bu, b, len(thu_tu) + 1))
        if isinstance(val(f"P{r}"), (int, float)):
            gv.append((ky, bu, b, val(f"P{r}")))
        vals = {ct: float(val(f"{c}{r}") or 0) for c, ct in COT_BAO_CAO.items()}
        dang_sd = vals["hsd_tren_12t"] + vals["hsd_6_12t"] + vals["hsd_duoi_6t"]
        clc = vals["clc_tren_24t"] + vals["clc_tren_12t"] + vals["clc_tren_06t"]
        for ct, v in (("tong_kd", dang_sd + vals["het_han"] + vals["cho_huy"] + vals["loi_kt"]),
                      ("dang_sd", dang_sd), ("het_han", vals["het_han"]), ("cho_huy", vals["cho_huy"]),
                      ("loi_kt", vals["loi_kt"]), ("clc", clc),
                      ("xau", vals["het_han"] + vals["cho_huy"] + vals["loi_kt"])):
            tkk.append((ky, bu, b, ct, v))
        for col, k in ky_truoc.items():
            if isinstance(val(f"{col}{r}"), (int, float)):
                tkk.append((k, bu, b, "tong_kd", val(f"{col}{r}")))

    # bảng cảnh báo tồn đọng mức khối (B..D = kỳ trước, E = kỳ này)
    for r in range(1, wv.max_row + 1):
        if norm(val(f"A{r}")).lower() == "chỉ tiêu":
            hdr = {c: parse_ngay(val(f"{c}{r}")) for c in "BCDE"
                   if isinstance(val(f"{c}{r}"), str) and re.match(r"\d{2}/\d{2}/\d{4}", val(f"{c}{r}"))}
            for rr in range(r + 1, r + 15):
                ct = CANH_BAO.get(norm(val(f"A{rr}")).lower())
                if not ct:
                    continue
                for c, k in hdr.items():
                    if k != ky and isinstance(val(f"{c}{rr}"), (int, float)):
                        tkk.append((k, "", "", ct, val(f"{c}{rr}")))
            break
    return thu_tu, gv, tkk


def build_sql(master, thu_tu, gv, tkk, khoi, ky, nguon):
    files = []
    hieu_luc = f"{ky.isoformat()} 00:00:00+07"
    cols = ("ma_bravo", "ten_mat_hang", "khoi", "bu", "nhom_sp", "hang", "pl_hang_hoa", "don_gia", "row_hash")
    for i in range(0, len(master), CHUNK):
        part = master[i:i + CHUNK]
        values = ",\n".join("(" + ",".join(q(m[c]) for c in cols) + ")" for m in part)
        files.append(f"""
with src({",".join(cols)}) as (values
{values}
),
cur as (
  select distinct on (d.ma_bravo) d.ma_bravo, d.row_hash
  from app_inventory.dm_hang_ton_kho d
  where d.ma_bravo in (select ma_bravo from src)
  order by d.ma_bravo, d.updated_at desc
)
insert into app_inventory.dm_hang_ton_kho ({",".join(cols)}, nguon, updated_at)
select {",".join("s." + c for c in cols)}, {q(nguon)}, {q(hieu_luc)}::timestamptz
from src s left join cur c on c.ma_bravo = s.ma_bravo
where c.ma_bravo is null or c.row_hash <> s.row_hash
on conflict (ma_bravo, updated_at) do update
  set ten_mat_hang = excluded.ten_mat_hang, khoi = excluded.khoi, bu = excluded.bu,
      nhom_sp = excluded.nhom_sp, hang = excluded.hang, pl_hang_hoa = excluded.pl_hang_hoa,
      don_gia = excluded.don_gia, row_hash = excluded.row_hash, nguon = excluded.nguon;
""")
    sql = []
    if thu_tu:
        sql.append("insert into app_inventory.dm_nhom_bao_cao (khoi, bu, nhom_sp, thu_tu) values\n"
                   + ",\n".join(f"({q(khoi)},{q(b)},{q(n)},{t})" for b, n, t in thu_tu)
                   + "\non conflict (khoi, bu, nhom_sp) do update set thu_tu = excluded.thu_tu;")
    if gv:
        sql.append("insert into app_inventory.gv_ban_tb (ky, bu, nhom_sp, gia_tri) values\n"
                   + ",\n".join(f"({q(k)},{q(b)},{q(n)},{v})" for k, b, n, v in gv)
                   + "\non conflict (ky, bu, nhom_sp) do update set gia_tri = excluded.gia_tri;")
    if tkk:
        sql.append("insert into app_inventory.ton_kho_ky (ky, khoi, bu, nhom_sp, chi_tieu, gia_tri, nguon) values\n"
                   + ",\n".join(f"({q(k)},{q(khoi)},{q(b)},{q(n)},{q(ct)},{v},{q('excel ' + nguon)})"
                                for k, b, n, ct, v in tkk)
                   + "\non conflict (ky, khoi, bu, nhom_sp, chi_tieu) do update"
                     " set gia_tri = excluded.gia_tri, nguon = excluded.nguon, updated_at = now();")
    files.append("\n\n".join(sql))
    return files


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--ky", help="ngày báo cáo YYYY-MM-DD (mặc định: theo tên file)")
    ap.add_argument("--sql-only", metavar="DIR", help="chỉ ghi file SQL ra thư mục, không chạy")
    args = ap.parse_args()

    ky = date.fromisoformat(args.ky) if args.ky else ky_tu_ten_file(args.file)
    nguon = Path(args.file).name
    print(f"Đọc {nguon} (kỳ {ky}) ...")
    wb_f = openpyxl.load_workbook(args.file, data_only=False)
    wb_v = openpyxl.load_workbook(args.file, data_only=True)

    # khối lấy từ dữ liệu gốc
    ws = wb_v[SHEET_RAW]
    hdr = [norm(c.value) for c in ws[1]]
    khoi_vals = {norm(r[hdr.index("Khối")]) for r in ws.iter_rows(min_row=2, values_only=True)} - {""}
    if len(khoi_vals) != 1:
        sys.exit(f"File cần đúng 1 khối, đang có: {khoi_vals}")
    khoi = khoi_vals.pop()

    thu_tu, gv, tkk = doc_bao_cao(wb_f, wb_v, khoi, ky)
    ten_nhom_chuan = {n.lower(): n for _, n, _ in thu_tu}
    master = doc_master(wb_v, ten_nhom_chuan)
    print(f"  master: {len(master)} mã | nhóm báo cáo: {len(thu_tu)} | GV bán TB: {len(gv)} | ton_kho_ky: {len(tkk)} dòng")

    files = build_sql(master, thu_tu, gv, tkk, khoi, ky, nguon)
    out_dir = Path(args.sql_only or os.path.join(os.path.dirname(__file__), "..", "out"))
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for i, s in enumerate(files):
        p = out_dir / f"nap_{ky.isoformat()}_{i:02d}.sql"
        p.write_text(s, encoding="utf-8")
        paths.append(p)
    if args.sql_only:
        print(f"  đã ghi {len(paths)} file SQL vào {out_dir}")
        return
    root = Path(__file__).resolve().parent.parent
    for p in paths:
        r = subprocess.run(["supabase", "db", "query", "--linked", "-f", str(p.resolve())],
                           cwd=root, capture_output=True, text=True, encoding="utf-8", shell=os.name == "nt")
        if r.returncode != 0 or '"error"' in r.stdout:
            sys.exit(f"Lỗi khi chạy {p.name}:\n{r.stdout}\n{r.stderr}")
        print(f"  ✓ {p.name}")
    print("Xong.")


if __name__ == "__main__":
    main()
